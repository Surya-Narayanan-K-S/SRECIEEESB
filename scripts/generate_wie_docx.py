import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

os.makedirs('public/downloads', exist_ok=True)
os.makedirs('src/assets', exist_ok=True)

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run('IEEE Women in Engineering (WIE)')
title_run.font.name = 'Calibri'
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(120, 35, 123) # Purple #78237b
title_p.paragraph_format.space_after = Pt(2)

# Subtitle
sub_p = doc.add_paragraph()
sub_run = sub_p.add_run('TESTIMONIALS & FEEDBACK REPORT (OFFICE BEARERS & MENTORS)')
sub_run.font.name = 'Calibri'
sub_run.font.size = Pt(13)
sub_run.font.bold = True
sub_run.font.color.rgb = RGBColor(0, 40, 85) # Navy #002855
sub_p.paragraph_format.space_after = Pt(14)

quotes = [
    {
        "quote": "“The WIE activities gave me the confidence to actively participate in STEM and explore new opportunities in engineering.”",
        "author": "Student Participant",
        "name": "G J Lithigaa",
        "role": "Chairperson, IEEE WIE (III IT)",
        "tag": "Confidence & Participation"
    },
    {
        "quote": "“The sessions were inspiring and helped me understand how technical knowledge can be combined with leadership and communication skills.”",
        "author": "Student Member",
        "name": "S Dhakshitha",
        "role": "Secretary, IEEE WIE (III CSE)",
        "tag": "Leadership & Technical Fusion"
    },
    {
        "quote": "“WIE creates a welcoming environment where students are encouraged to learn, collaborate and grow together.”",
        "author": "WIE Member",
        "name": "S Tejasvi",
        "role": "Joint Activity Coordinator (III BME)",
        "tag": "Community & Collaboration"
    },
    {
        "quote": "“The programs provided valuable exposure to real-world engineering, professional development and emerging technologies.”",
        "author": "Event Participant",
        "name": "R Tejashri",
        "role": "Treasurer, IEEE WIE (III BME)",
        "tag": "Real-World Exposure"
    },
    {
        "quote": "“The mentorship and interactions encouraged me to think beyond academics and consider new possibilities in my engineering career.”",
        "author": "Student Participant",
        "name": "J Sindhu",
        "role": "Social Media Lead (III M.Tech CSE)",
        "tag": "Mentorship & Career Vision"
    },
    {
        "quote": "“WIE is more than a platform for events—it is a community that inspires women to learn, lead and contribute to STEM.”",
        "author": "Mentor / Volunteer",
        "name": "Mrs. S. Jansi Rani",
        "role": "IEEE WIE Coordinator & Faculty Advisor · AP (Sr.G)/IT",
        "tag": "Empowerment & STEM Impact"
    }
]

for idx, item in enumerate(quotes, 1):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    
    # Border & padding via XML
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), '78237B')
    tcBorders.append(left)
    
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'none')
    tcBorders.append(top)
    
    right = OxmlElement('w:right')
    right.set(qn('w:val'), 'none')
    tcBorders.append(right)
    
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'none')
    tcBorders.append(bottom)
    
    tcPr.append(tcBorders)
    
    # Shading (light purple/slate bg)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F9F5FA')
    tcPr.append(shd)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.left_indent = Inches(0.15)
    cp.paragraph_format.right_indent = Inches(0.15)
    
    q_run = cp.add_run(item['quote'])
    q_run.font.name = 'Georgia'
    q_run.font.size = Pt(11)
    q_run.font.bold = True
    q_run.font.color.rgb = RGBColor(30, 41, 59)
    
    author_p = cell.add_paragraph()
    author_p.paragraph_format.space_before = Pt(2)
    author_p.paragraph_format.space_after = Pt(6)
    author_p.paragraph_format.left_indent = Inches(0.15)
    author_p.paragraph_format.right_indent = Inches(0.15)
    
    dash_run = author_p.add_run('— ')
    dash_run.font.color.rgb = RGBColor(120, 35, 123)
    
    name_run = author_p.add_run(f"{item['name']} ")
    name_run.font.name = 'Calibri'
    name_run.font.size = Pt(10.5)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(120, 35, 123)
    
    role_run = author_p.add_run(f"({item['role']} · {item['author']})")
    role_run.font.name = 'Calibri'
    role_run.font.size = Pt(9.5)
    role_run.font.italic = True
    role_run.font.color.rgb = RGBColor(71, 85, 105)
    
    tag_run = author_p.add_run(f"   [{item['tag']}]")
    tag_run.font.name = 'Calibri'
    tag_run.font.size = Pt(9)
    tag_run.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# Note Section
note_p = doc.add_paragraph()
note_p.paragraph_format.space_before = Pt(12)
note_run_label = note_p.add_run('Note: ')
note_run_label.font.name = 'Calibri'
note_run_label.font.size = Pt(10)
note_run_label.font.bold = True
note_run_label.font.color.rgb = RGBColor(180, 83, 9) # Amber

note_run_text = note_p.add_run('These are sample presentation-ready quotes paired with official IEEE WIE office bearers and faculty mentor profiles. For an official report, replace them with actual signed feedback collected from your WIE members, participants, or mentors.')
note_run_text.font.name = 'Calibri'
note_run_text.font.size = Pt(9.5)
note_run_text.font.italic = True
note_run_text.font.color.rgb = RGBColor(100, 116, 139)

doc.save('public/downloads/WIE_Testimonials_Feedback.docx')
doc.save('src/assets/WIE_Testimonials_Feedback.docx')
print('Updated DOCX successfully created at public/downloads/WIE_Testimonials_Feedback.docx')
